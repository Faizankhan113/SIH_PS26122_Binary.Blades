from __future__ import annotations

import csv
import io
import json
import os
import secrets
import sys
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from functools import wraps
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from flask import Flask, abort, g, jsonify, redirect, render_template, request, session, url_for
from werkzeug.utils import secure_filename

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src import pipeline_runs  # noqa: E402
from src.timeutil import (  # noqa: E402  (needs ROOT on sys.path first)
    actual_day,
    canonical_actual,
    canonicalize_event_dates,
    duration_hours,
    format_actual,
    today_local,
)

# Read .env BEFORE the settings below (FLASK_SECRET_KEY, FLASK_DEBUG, ...).
load_dotenv(ROOT / ".env")


def _load_secret_key() -> tuple[str, bool]:
    """The signing key comes from FLASK_SECRET_KEY; there is no built-in default.

    Returns (key, generated). When the variable is missing a random key is made
    for this process only, so nobody can forge a session cookie with a key that
    is written in the source code.
    """
    key = (os.getenv("FLASK_SECRET_KEY") or "").strip()
    if key:
        return key, False
    return secrets.token_hex(32), True


def _debug_enabled() -> bool:
    """Flask debug mode (auto-reload, interactive debugger) only when FLASK_DEBUG=1."""
    return os.getenv("FLASK_DEBUG", "").strip() == "1"


app = Flask(__name__)
app.secret_key, _SECRET_KEY_GENERATED = _load_secret_key()
if _SECRET_KEY_GENERATED:
    _msg = (
        "FLASK_SECRET_KEY is not set: using a random key for this run only. Everyone is logged "
        "out whenever the server restarts. Set FLASK_SECRET_KEY (see .env.example) to keep logins."
    )
    print(f"WARNING: {_msg}", file=sys.stderr)
    app.logger.warning(_msg)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024


# --------------------------------------------------------------------------
# Auth: route guards
#
# The session cookie holds just {user_id, user_name, role} and, while a report
# is being processed, the run_id of the current run. The statements,
# extraction results and match rows live on the server (src/pipeline_runs.py),
# so the cookie stays far below the browser limit of about 4 KB.
# --------------------------------------------------------------------------

def current_user() -> dict[str, Any] | None:
    if not session.get("user_id"):
        return None
    return {"user_id": session["user_id"], "name": session["user_name"], "role": session["role"]}


def _unauthorized_response():
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": "Please log in first."}), 401
    return redirect(url_for("login", next=request.path))


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return _unauthorized_response()
        return view(*args, **kwargs)
    return wrapped


def role_required(*roles: str):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not session.get("user_id"):
                return _unauthorized_response()
            if session.get("role") not in roles:
                if request.path.startswith("/api/"):
                    return jsonify({"ok": False, "error": "Not authorized for this role."}), 403
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


@app.context_processor
def inject_current_user():
    user = current_user()
    open_planner_items = None
    if user and user["role"] == "supervisor":
        # The number shown next to the Planner List link. A database problem
        # must never break the page that is being rendered, so it just hides the number.
        try:
            from src.database import count_open_planner_items
            open_planner_items = count_open_planner_items()
        except Exception:
            app.logger.warning("Could not count open planner items", exc_info=True)
    return {"current_user": user, "open_planner_items": open_planner_items}


@app.template_filter("fmt_actual")
def fmt_actual_filter(value: Any, time_stated: Any = None) -> str:
    """An actual start/finish for display. The time of day appears ONLY when
    the report stated one (never a fake 00:00). Empty string when there is no value."""
    return format_actual(value, time_stated)


@app.template_global("actual_hours")
def actual_hours(start: Any, start_stated: Any, finish: Any, finish_stated: Any) -> float | None:
    """Actual duration in hours, only when BOTH times were stated."""
    return duration_hours(start, start_stated, finish, finish_stated)


@app.errorhandler(403)
def forbidden(_exc):
    return render_template("error.html", step=0, hide_stepper=True, message="You don't have access to this page."), 403

UPLOAD_DIR = ROOT / "data" / "ui_uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

# Spreadsheet columns. Statement text: the first of these headers found.
_STATEMENT_COLUMNS = ("statement", "description", "update", "progress")
# A column holding a schedule identifier, in order of preference.
_ID_COLUMNS = ("activity code", "activity id", "l6", "wbs")


def _header_key(cell: Any) -> str:
    """'Activity_ID ' -> 'activity id'."""
    return " ".join(str(cell or "").replace("_", " ").lower().split())


def _cell_text(row: Any, index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _attach_activity_ref(statement: str, ref: str) -> str:
    """Write the row's activity ID into the statement text.

    Statements travel through the whole pipeline as plain text, so putting the ID
    in the text is enough for both the extraction prompt (rule 13) and the
    matcher (which also scans the raw text for exact codes) to see it.
    """
    ref = ref.strip()
    if not ref or ref.upper() in statement.upper():
        return statement
    return f"{statement} [Activity ref: {ref}]"


def _statements_from_rows(rows: list[Any], fallback: str) -> list[str]:
    """Turn spreadsheet rows (first row = header) into statements.

    * The statement is the `statement`/`description`/`update`/`progress` column.
      With none of them, `fallback` picks the "last" or the "first" column that
      is not the ID column, so an ID column is never mistaken for the text.
    * A column named `activity code`, `activity id`, `l6` or `wbs` is attached to
      each statement.
    """
    if not rows:
        return []
    header = [_header_key(h) for h in rows[0]]
    id_idx = next((header.index(name) for name in _ID_COLUMNS if name in header), None)
    named = [header.index(name) for name in _STATEMENT_COLUMNS if name in header]
    if named:
        idx = named[0]
    else:
        usable = [i for i in range(len(header)) if i != id_idx] or [0]
        idx = usable[-1] if fallback == "last" else usable[0]
    data_rows = rows[1:] if any(header) else rows
    out = []
    for row in data_rows:
        statement = _cell_text(row, idx)
        if not statement:
            continue
        if id_idx is not None and id_idx != idx:
            statement = _attach_activity_ref(statement, _cell_text(row, id_idx))
        out.append(statement)
    return out


def _read_uploaded_file(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".log"}:
        return [x.strip() for x in path.read_text(encoding="utf-8", errors="ignore").splitlines() if x.strip()]
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8", errors="ignore")
        # Prefer a column named statement / description / update; otherwise the last column.
        return _statements_from_rows(list(csv.reader(io.StringIO(text))), fallback="last")
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        return _statements_from_rows(list(ws.iter_rows(values_only=True)), fallback="first")
    if suffix == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF support needs PyMuPDF (fitz).") from exc
        doc = fitz.open(path)
        text = "\n".join(page.get_text() for page in doc)
        return [x.strip() for x in text.splitlines() if x.strip()]
    raise ValueError(f"Unsupported file type: {suffix}")


def _keyword_fallback_extract(statement: str, report_date: str) -> dict[str, Any]:
    """Keyword-rule fallback used only when LLM extraction fails. Its output is always flagged as degraded."""
    lower = statement.lower()
    event = {
        "event_id": f"EVT-{uuid.uuid4().hex[:8].upper()}",
        "raw_text": statement,
        "activity_description": "spool erection" if "spool" in lower else "work activity",
        "asset": None,
        "discipline": "Piping" if any(x in lower for x in ("piping", "spool", "valve", "erection")) else None,
        "location": None,
        "status": "COMPLETED" if any(x in lower for x in ("completed", "finished", "done")) else "IN_PROGRESS",
        "actual_start": None,
        "actual_finish": None,
        # A degraded (keyword-guess) event never carries a time of day.
        "actual_start_time_stated": False,
        "actual_finish_time_stated": False,
        "progress_percent": None,
        "contractor": None,
        "delay_reason_reported": None,
        "evidence_quote": statement,
        "extraction_degraded": True,
        "extraction_error": None,
    }
    import re
    m = re.search(r"line\s*[- ]?(\d+)", statement, flags=re.I)
    if m:
        event["asset"] = f"Line {m.group(1)}"
    m = re.search(r"valve\s*[- ]?([A-Za-z0-9-]+)", statement, flags=re.I)
    if m:
        event["asset"] = f"Valve {m.group(1)}"
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:%|percent)", statement, flags=re.I)
    if m:
        event["progress_percent"] = float(m.group(1))
        event["status"] = "COMPLETED" if event["progress_percent"] >= 100 else "IN_PROGRESS"
    if "today" in lower:
        event["actual_finish"] = report_date if event["status"] == "COMPLETED" else None
    return event


def initialize_db() -> None:
    """Create/migrate the database and sync the schedule -- ONCE, at app start.

    Nothing else in the request path initializes the database any more, so if
    this fails the app must be restarted after fixing the cause.
    """
    try:
        from src.database import initialize_database
        summary = initialize_database()
        app.logger.info("Database ready (schedule sync: %s)", summary)
        # Drop old scratch runs, and fail extraction items a previous
        # server process left half-done (their worker threads no longer exist).
        removed = pipeline_runs.cleanup_old_runs()
        interrupted = pipeline_runs.recover_interrupted_items()
        app.logger.info("Pipeline runs: %s old run(s) removed, %s interrupted item(s) marked failed", removed, interrupted)
    except Exception:
        app.logger.exception("Database initialization failed at startup; fix the cause and restart")


def _save_event(event: dict[str, Any], match: dict[str, Any], ingestion: dict[str, Any]) -> tuple[dict[str, Any] | None, str | None]:
    """Save one matched event; return (result, None) or (None, error message).

    All plan-safety rules live inside `save_progress_update`;
    nothing is re-implemented here.
    """
    try:
        from src.database import save_progress_update
        return save_progress_update(
            event=event,
            match=match,
            reported_by=ingestion.get("reported_by", "Unknown Reporter"),
            source_name=ingestion.get("source_name", "Pasted text"),
            report_date=ingestion.get("report_date"),
            report_id=ingestion.get("report_id"),
            reported_by_user_id=ingestion.get("reported_by_user_id"),
        ), None
    except Exception as exc:
        app.logger.warning("Database persistence failed: %s", exc)
        return None, str(exc) or exc.__class__.__name__


def save_update_to_db(event: dict[str, Any], match: dict[str, Any], ingestion: dict[str, Any]) -> dict[str, Any] | None:
    """Save one matched event. Returns the save result, or None if it failed."""
    return _save_event(event, match, ingestion)[0]


initialize_db()


def _needs_correction_payload(statement: str, exc: Exception) -> dict[str, Any]:
    """A statement whose CONTENT is invalid (for example finish before start).

    Not an outage, so no keyword guess is made. The item carries the reason and
    the original text; it is shown to the user and is never matched.
    """
    return {
        "event_id": f"EVT-{uuid.uuid4().hex[:8].upper()}",
        "raw_text": statement,
        "needs_correction": True,
        "validation_error": str(exc),
        "extraction_degraded": False,
        "extraction_error": None,
    }


def real_extract(statement: str, report_date: str, message_time: datetime | None = None) -> dict[str, Any]:
    """Extract one statement.

    Three outcomes:
      * a normal event dict,
      * a "needs correction" dict (needs_correction=True) when the statement
        itself is invalid -- no guessing,
      * a degraded keyword-guess event (extraction_degraded=True) ONLY when the
        provider/network failed.
    """
    from src.normalize import ExtractionValidationError

    try:
        from src.extract import extract_progress_event
        # `message_time` (a chat message's own timestamp) is passed only when
        # there is one; pasted reports and files have no such moment.
        extra = {"message_time": message_time} if message_time is not None else {}
        result = extract_progress_event(statement, report_date=date.fromisoformat(report_date), **extra)
        # A date-only value must not leave here as "...T00:00:00".
        payload = canonicalize_event_dates(result.model_dump(mode="json"))
        payload.setdefault("extraction_degraded", False)
        payload.setdefault("extraction_error", None)
        return payload
    except ExtractionValidationError as exc:
        app.logger.info("Statement needs correction (%s): %s", exc, statement)
        return _needs_correction_payload(statement, exc)
    except Exception as exc:
        app.logger.exception("LLM extraction failed for statement, falling back to keyword heuristics")
        event = _keyword_fallback_extract(statement, report_date)
        # Distinct, visible error state -- the UI must show this was a
        # degraded keyword-rule guess, not trustworthy LLM extraction, so a
        # provider outage doesn't silently feed wrong data into matching.
        # The flag is carried through matching and can never auto-accept.
        event["extraction_degraded"] = True
        event["extraction_error"] = f"{type(exc).__name__}: {exc}"
        return event


class ScheduleUnavailableError(RuntimeError):
    """Raised when the schedule can't be loaded from the database. Callers must
    surface this as an error state and never match against a partial schedule."""


def get_match_schedule() -> list[dict[str, Any]]:
    """Active (non-archived) planned activities the matcher may choose from.

    Call this ONCE per batch of events and pass the result to `match_event`.
    """
    try:
        from src.database import list_active_schedule
        return list_active_schedule()
    except Exception as exc:
        app.logger.exception("Schedule database unavailable")
        raise ScheduleUnavailableError(
            "The schedule could not be loaded from the database, so matching was not run. "
            "Check the PostgreSQL connection (DATABASE_URL) and try again."
        ) from exc


def guard_degraded_match(event: dict[str, Any], match: dict[str, Any]) -> dict[str, Any]:
    """A degraded event can never be AUTO_ACCEPTED.

    Always copies the degraded flag and error onto the match (so the matched
    screen, the history row and the supervisor queue can show the badge).
    An AUTO_ACCEPTED result is forced to REVIEW with the reason prepended.
    REVIEW stays REVIEW. UNMATCHED and ERROR keep their status: they have no
    target activity, so there is nothing for a supervisor to approve, and they
    write nothing to the plan anyway.
    """
    degraded = bool(event.get("extraction_degraded"))
    match["extraction_degraded"] = degraded
    match["extraction_error"] = event.get("extraction_error") if degraded else None
    if not degraded:
        return match
    status = match.get("review_status")
    if status == "AUTO_ACCEPTED":
        match["review_status"] = "REVIEW"
        status = "REVIEW"
    if status in {"REVIEW", "UNMATCHED"}:
        from src.database import DEGRADED_REASON
        reason = match.get("reason") or ""
        if DEGRADED_REASON not in reason:
            match["reason"] = f"{DEGRADED_REASON} {reason}".strip()
    return match


def _schedule_unavailable_match(event: dict[str, Any], exc: Exception) -> dict[str, Any]:
    return guard_degraded_match(event, {
        "event_id": event.get("event_id"),
        "matched_activity_id": None,
        "match_method": "ERROR",
        "confidence": 0.0,
        "reason": f"Schedule unavailable, matching not attempted: {exc}",
        "review_status": "ERROR",
        "candidates": [],
    })


def match_event(event: dict[str, Any], schedule: list[dict[str, Any]]) -> dict[str, Any]:
    """UI adapter for the hybrid retrieval + LLM-rerank matcher.

    `schedule` is the already-loaded active schedule (see `get_match_schedule`);
    it is loaded once per batch, not once per event.

    Works directly on plain dicts end to end (event from real_extract(),
    schedule rows from the DB) -- no Pydantic <-> dict mismatch, so this
    actually exercises the real matching pipeline instead of silently
    falling back to a crude heuristic.
    """
    try:
        from src.semantic_match import match_progress_event_v2
        return guard_degraded_match(event, dict(match_progress_event_v2(event, schedule)))
    except Exception as exc:
        # Fail loudly in the logs (never silently), but keep the UI usable:
        # surface this as an honest error state rather than a fabricated match.
        app.logger.exception("Matching engine failed for event %s", event.get("event_id"))
        return guard_degraded_match(event, {
            "event_id": event.get("event_id"),
            "matched_activity_id": None,
            "match_method": "MATCH_ENGINE_ERROR",
            "confidence": 0.0,
            # A processing failure is ERROR, never UNMATCHED. UNMATCHED means
            # "the matcher ran and found no activity" (a possible new activity).
            "reason": f"Matching engine raised an error and did not produce a result: {exc}",
            "review_status": "ERROR",
            "candidates": [],
        })


def _before_after(event: dict[str, Any], match: dict[str, Any]) -> dict[str, Any]:
    before = {
        "event_id": event.get("event_id"),
        "activity_description": event.get("activity_description"),
        "asset": event.get("asset"),
        "discipline": event.get("discipline"),
        "status": event.get("status"),
        "actual_start": event.get("actual_start"),
        "actual_finish": event.get("actual_finish"),
        "actual_start_time_stated": event.get("actual_start_time_stated", False),
        "actual_finish_time_stated": event.get("actual_finish_time_stated", False),
        "progress_percent": event.get("progress_percent"),
        "contractor": event.get("contractor"),
    }
    after = dict(before)
    after.update(
        {
            "matched_activity_id": match.get("matched_activity_id"),
            "match_method": match.get("match_method"),
            "confidence": match.get("confidence"),
            "review_status": match.get("review_status"),
        }
    )
    return {"before": before, "after": after}


@app.route("/")
def index():
    if not session.get("user_id"):
        return redirect(url_for("login"))
    if session.get("role") == "supervisor":
        return redirect(url_for("review"))
    return redirect(url_for("ingest"))


# --------------------------------------------------------------------------
# Auth routes
# --------------------------------------------------------------------------

@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("index"))
    error = None
    if request.method == "POST":
        from src.auth import verify_login
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        user, error = verify_login(username, password)
        if user:
            session.clear()
            session["user_id"] = user["user_id"]
            session["user_name"] = user["name"]
            session["role"] = user["role"]
            next_url = request.args.get("next") or request.form.get("next")
            if next_url and next_url.startswith("/"):
                return redirect(next_url)
            return redirect(url_for("review" if user["role"] == "supervisor" else "ingest"))
    return render_template("login.html", step=0, hide_stepper=True, error=error, next=request.args.get("next", ""))


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if session.get("user_id"):
        return redirect(url_for("index"))
    error = None
    submitted = False
    if request.method == "POST":
        from src.auth import UsernameTakenError, signup_contractor
        name = request.form.get("name", "")
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")
        if password != confirm:
            error = "Passwords don't match."
        else:
            try:
                signup_contractor(name=name, username=username, password=password)
                submitted = True
            except (ValueError, UsernameTakenError) as exc:
                error = str(exc)
    return render_template("signup.html", step=0, hide_stepper=True, error=error, submitted=submitted)


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# --------------------------------------------------------------------------
# Supervisor: review queue
# --------------------------------------------------------------------------

@app.route("/review")
@role_required("supervisor")
def review():
    from src.auth import list_pending_contractor_accounts
    from src.database import list_review_updates

    try:
        pending_reviews = list_review_updates()
        db_error = None
    except Exception as exc:
        app.logger.exception("Failed to load review queue")
        pending_reviews, db_error = [], str(exc)

    return render_template(
        "review.html",
        step=0,
        hide_stepper=True,
        pending_accounts=list_pending_contractor_accounts(),
        pending_reviews=pending_reviews,
        db_error=db_error,
    )


@app.post("/api/accounts/<user_id>/approve")
@role_required("supervisor")
def api_approve_account(user_id: str):
    from src.auth import STATUS_ACTIVE, set_account_status
    set_account_status(user_id, STATUS_ACTIVE)
    return jsonify({"ok": True})


@app.post("/api/accounts/<user_id>/reject")
@role_required("supervisor")
def api_reject_account(user_id: str):
    from src.auth import STATUS_REJECTED, set_account_status
    set_account_status(user_id, STATUS_REJECTED)
    return jsonify({"ok": True})


def _not_awaiting_review_response(update_id: str):
    """Say WHY there is nothing to decide instead of one vague message.

    Unknown id -> 404. An item that was already approved or rejected -> 409 with
    who decided it, so a second click (or a second supervisor) gets a clear
    message and never a second update.
    """
    from src.database import get_progress_update
    row = get_progress_update(update_id)
    if row is None:
        return jsonify({"ok": False, "error": "That update does not exist."}), 404
    status = row.get("match_status")
    if status == "REJECTED":
        message = "This update was already rejected. Nothing was changed."
    elif status == "REVIEW":
        # Still pending, but it has no target activity to apply to.
        return jsonify({"ok": False, "error": "This update has no matched activity, so it cannot be approved."}), 409
    else:
        who = row.get("approved_by")
        message = "This update was already approved" + (f" by {who}" if who and who != "SYSTEM" else "") + \
                  ". It was applied once; nothing was changed."
    return jsonify({"ok": False, "error": message, "already_decided": True, "match_status": status}), 409


@app.post("/api/reviews/<update_id>/approve")
@role_required("supervisor")
def api_approve_review(update_id: str):
    from src.database import DegradedConfirmationRequired, decide_review_update
    body = request.get_json(silent=True) or {}
    try:
        result = decide_review_update(
            update_id,
            decision="approve",
            supervisor=current_user(),
            confirm_degraded=body.get("confirm_degraded") is True,
        )
    except DegradedConfirmationRequired as exc:
        # The UI shows the warning and re-sends with confirm_degraded=true.
        return jsonify({"ok": False, "error": str(exc), "needs_confirmation": True}), 409
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 409
    if not result:
        return _not_awaiting_review_response(update_id)
    return jsonify({"ok": True, **result})


@app.post("/api/reviews/<update_id>/reject")
@role_required("supervisor")
def api_reject_review(update_id: str):
    from src.database import decide_review_update
    try:
        result = decide_review_update(update_id, decision="reject", supervisor=current_user())
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 409
    if not result:
        return _not_awaiting_review_response(update_id)
    return jsonify({"ok": True, **result})


# --------------------------------------------------------------------------
# Supervisor: planner list
#
# UNMATCHED (NOT_FOUND) statements are flagged here for a planner. Nothing on this page
# edits the schedule or the plan: "Mark as new activity" only records a note.
# --------------------------------------------------------------------------

PLANNER_NOTE_MAX = 500


@app.route("/planner")
@role_required("supervisor")
def planner():
    from src.database import PLANNER_STATUSES, list_planner_items, planner_status_counts

    status = (request.args.get("status") or "open").strip().lower()
    if status not in (*PLANNER_STATUSES, "all"):
        status = "open"
    try:
        items = list_planner_items(status)
        counts = planner_status_counts()
        db_error = None
    except Exception as exc:
        app.logger.exception("Failed to load the planner list")
        items, counts, db_error = [], {}, str(exc)
    return render_template(
        "planner.html",
        step=0,
        hide_stepper=True,
        items=items,
        counts=counts,
        status=status,
        db_error=db_error,
    )


def _planner_decision_response(update_id: str, decision: str):
    """Shared by the two planner actions."""
    from src.database import PlannerItemNotFound, decide_planner_item

    body = request.get_json(silent=True) or {}
    note = body.get("note")
    if note is not None and not isinstance(note, str):
        return jsonify({"ok": False, "error": "The note must be text."}), 400
    if note and len(note.strip()) > PLANNER_NOTE_MAX:
        return jsonify({"ok": False, "error": f"The note is too long (at most {PLANNER_NOTE_MAX} characters)."}), 400
    try:
        result = decide_planner_item(update_id, decision=decision, planner=current_user(), note=note)
    except PlannerItemNotFound:
        return jsonify({"ok": False, "error": "That item is not on the planner list."}), 404
    if not result["changed"]:
        # Already decided (an earlier click, or another supervisor): keep the first decision.
        label = "dismissed" if result["planner_status"] == "dismissed" else "marked as a new activity"
        who = result.get("planner_decided_by")
        return jsonify({
            "ok": False,
            "already_decided": True,
            "planner_status": result["planner_status"],
            "error": f"This item was already {label}" + (f" by {who}" if who else "") +
                     ". The first decision was kept; nothing was changed.",
        }), 409
    return jsonify({"ok": True, **result})


@app.post("/api/planner/<update_id>/dismiss")
@role_required("supervisor")
def api_planner_dismiss(update_id: str):
    return _planner_decision_response(update_id, "dismiss")


@app.post("/api/planner/<update_id>/convert")
@role_required("supervisor")
def api_planner_convert(update_id: str):
    return _planner_decision_response(update_id, "convert")


# --------------------------------------------------------------------------
# Contractor: existing ingest -> extract -> match flow, now gated
# --------------------------------------------------------------------------

# Session keys older versions used to store whole result sets in the cookie.
_LEGACY_SESSION_KEYS = ("ingestion", "extracted_results", "match_results", "db_updates", "match_db_rows", "before_after")


def _current_run() -> dict[str, Any] | None:
    """The logged-in user's current run, or None.

    Looks the run up by the id in the session AND the user id, so a run id that
    belongs to someone else (or no longer exists) simply yields no run.
    """
    if "current_run" not in g:
        user = current_user()
        run_id = session.get("run_id")
        g.current_run = pipeline_runs.get_run(run_id, user["user_id"]) if user and run_id else None
        if run_id and g.current_run is None:
            session.pop("run_id", None)  # stale id: expired, deleted or someone else's
    return g.current_run


def _stepper_availability() -> dict[str, bool]:
    """Return navigation availability for the current pipeline run."""
    run = _current_run()
    return {"can_extract": bool(run), "can_match": bool(run and run["has_match"])}


@app.route("/ingest")
@role_required("contractor")
def ingest():
    run = _current_run()
    return render_template(
        "ingest.html",
        step=1,
        ingestion=run["ingestion"] if run else None,
        **_stepper_availability(),
    )


@app.post("/api/ingest")
@role_required("contractor")
def api_ingest():
    report_date = request.form.get("report_date") or today_local().isoformat()
    # "Reported by" is locked to the logged-in contractor, not a free-text
    # field anyone can type into.
    user = current_user()
    reported_by = user["name"]
    reported_by_user_id = user["user_id"]
    # Every statement of one submission must land in the same report row, so a
    # report id always exists (typed by the contractor, or generated here).
    report_id = (request.form.get("report_id") or "").strip() or f"RPT-{uuid.uuid4().hex[:10].upper()}"
    statements: list[str] = []
    source_name = "Pasted text"

    raw_text = request.form.get("raw_text", "").strip()
    uploaded = request.files.get("file")
    if uploaded and uploaded.filename:
        filename = secure_filename(uploaded.filename)
        path = UPLOAD_DIR / filename
        uploaded.save(path)
        source_name = filename
        statements = _read_uploaded_file(path)
    elif raw_text:
        statements = [x.strip() for x in raw_text.splitlines() if x.strip()]

    if not statements:
        return jsonify({"ok": False, "error": "Upload a file or paste at least one progress statement."}), 400

    # A new ingestion starts a new run; the previous one (and anything the
    # stepper could unlock from it) is discarded.
    previous = _current_run()
    run_id = pipeline_runs.create_run(user["user_id"], {
        "source_name": source_name,
        "report_date": report_date,
        "reported_by": reported_by,
        "reported_by_user_id": reported_by_user_id,
        "report_id": report_id,
        "statements": statements,
    })
    if previous:
        pipeline_runs.delete_run(previous["run_id"])
    for key in _LEGACY_SESSION_KEYS:
        session.pop(key, None)
    session["run_id"] = run_id
    g.pop("current_run", None)
    try:
        pipeline_runs.cleanup_old_runs()
    except Exception:
        app.logger.warning("Cleanup of old pipeline runs failed", exc_info=True)
    return jsonify({"ok": True, "redirect": url_for("extracted")})


# --------------------------------------------------------------------------
# Extraction runs in the background, statement by statement
# --------------------------------------------------------------------------

def _worker_count() -> int:
    try:
        return max(1, int(os.getenv("PS26122_EXTRACT_WORKERS", "").strip() or 4))
    except ValueError:
        return 4


# One small pool for the whole process: at most this many statements are sent
# to the LLM at the same time, however many people are working. Together with
# the retry/backoff in src/providers.py this keeps us inside provider rate limits.
EXTRACT_WORKERS = _worker_count()
_EXTRACT_POOL = ThreadPoolExecutor(max_workers=EXTRACT_WORKERS, thread_name_prefix="extract")


def _extract_one(run_id: str, idx: int, statement: str, report_date: str) -> None:
    """Worker: extract one statement and save the outcome into its run row.

    Whatever happens, the item ends as 'done' or 'failed', so one bad statement
    can never leave the screen waiting forever or hold up the others.
    """
    try:
        result = real_extract(statement, report_date)
        pipeline_runs.save_item_result(run_id, idx, result)
    except Exception as exc:
        app.logger.exception("Extraction failed for statement %s of run %s", idx, run_id)
        try:
            pipeline_runs.save_item_failure(run_id, idx, f"{type(exc).__name__}: {exc}")
        except Exception:
            app.logger.exception("Could not record the failure of statement %s of run %s", idx, run_id)


def _start_extraction(run: dict[str, Any]) -> int:
    """Queue every pending (or previously failed) statement of the run. Returns how many were queued."""
    report_date = run["ingestion"]["report_date"]
    queued = 0
    for item in pipeline_runs.list_items(run["run_id"]):
        if item["state"] in (pipeline_runs.STATE_PENDING, pipeline_runs.STATE_FAILED) and pipeline_runs.claim_item(
            run["run_id"], item["idx"]
        ):
            _EXTRACT_POOL.submit(_extract_one, run["run_id"], item["idx"], item["statement"], report_date)
            queued += 1
    return queued


def _progress_payload(items: list[dict[str, Any]]) -> dict[str, Any]:
    progress = pipeline_runs.extraction_progress(items)
    progress["items"] = [
        {
            "idx": i["idx"],
            "state": i["state"],
            "error": i["error"],
            "degraded": bool(i["result"] and i["result"].get("extraction_degraded")),
            "needs_correction": bool(i["result"] and i["result"].get("needs_correction")),
        }
        for i in items
    ]
    return progress


@app.post("/api/extract/start")
@role_required("contractor")
def api_extract_start():
    """Explicit action: begin (or retry failed) extraction for the current run. Safe to call twice."""
    run = _current_run()
    if not run:
        return jsonify({"ok": False, "error": "No report is being processed. Start from the ingestion screen."}), 400
    queued = _start_extraction(run)
    return jsonify({"ok": True, "queued": queued, **_progress_payload(pipeline_runs.list_items(run["run_id"]))})


@app.get("/api/extract/status")
@role_required("contractor")
def api_extract_status():
    run = _current_run()
    if not run:
        return jsonify({"ok": False, "error": "No report is being processed."}), 404
    return jsonify({"ok": True, **_progress_payload(pipeline_runs.list_items(run["run_id"]))})


@app.route("/extracted")
@role_required("contractor")
def extracted():
    run = _current_run()
    if not run:
        return redirect(url_for("ingest"))
    ingestion = run["ingestion"]
    items = pipeline_runs.list_items(run["run_id"])
    progress = pipeline_runs.extraction_progress(items)
    common = dict(
        step=2,
        source=ingestion["source_name"],
        report_date=ingestion["report_date"],
        items=items,
        progress=progress,
        **_stepper_availability(),
    )
    if not progress["complete"]:
        # Loading screen: its script calls /api/extract/start, then polls /api/extract/status.
        return render_template("extracting.html", **common)
    return render_template("extracted.html", **common)


def _build_match_row(
    event: dict[str, Any],
    match: dict[str, Any],
    ingestion: dict[str, Any],
    get_planned_activity,
) -> dict[str, Any]:
    """One row of the matched screen. This SAVES the event, then reads the plan back.

    Called only from `continue_to_match` and `retry_match`; opening or reloading
    /matched only reads the stored row and never saves.

    * AUTO_ACCEPTED -> the plan is updated (unless the save downgrades it, or the
                       report is outdated / changes nothing: HISTORICAL / NO_CHANGE).
    * REVIEW        -> saved as pending, plan unchanged.
    * UNMATCHED     -> saved (as NOT_FOUND) for the planner list.
    * ERROR         -> saved as ERROR for the audit trail; never queued anywhere.
                       Retrying adds a new attempt to the same statement.

    `before_db` is read just before the save and `after_db` just after it, both
    from the database, so what the screen shows is what the database holds. The
    status shown is the status actually stored (a conflict turns AUTO_ACCEPTED
    into REVIEW).
    """
    status = match.get("review_status")
    target = match.get("matched_activity_id") if status in ("AUTO_ACCEPTED", "REVIEW") else None

    def read(l6_id):
        return get_planned_activity(l6_id) if (get_planned_activity and l6_id) else None

    before_db = read(target)
    saved: dict[str, Any] | None = None
    save_error: str | None = None
    outcome = "ERROR"
    if status in ("AUTO_ACCEPTED", "REVIEW", "UNMATCHED", "ERROR"):
        saved, save_error = _save_event(event, match, ingestion)
        if saved is None:
            outcome = "NOT_SAVED"
        else:
            if saved["duplicate"]:
                outcome = "DUPLICATE"
            else:
                # Show the status that was really stored (a conflict or a degraded
                # extraction can turn AUTO_ACCEPTED into REVIEW).
                match["review_status"] = saved["match_status"]
                match["reason"] = saved["match_reason"]
                outcome = {"AUTO_ACCEPTED": "APPLIED", "REVIEW": "PENDING",
                           "UNMATCHED": "LOGGED", "ERROR": "ERROR"}.get(saved["match_status"], "LOGGED")
                if saved["match_status"] == "AUTO_ACCEPTED" and not saved["applied"]:
                    # Accepted, but the live state was not changed: an older report
                    # than the one already applied, or one that agrees with it.
                    outcome = "HISTORICAL" if saved.get("outcome") == "HISTORICAL_ONLY" else "NO_CHANGE"
                if saved["conflicts"]:
                    outcome = "CONFLICT"
                    match["conflicts"] = saved["conflicts"]
                if saved["match_status"] == "REVIEW":
                    target = saved["l6_id"] or target
                    before_db = before_db or read(target)
            match["update_id"] = saved["update_id"]

    after_db = read((saved or {}).get("l6_id") or target) if target or (saved or {}).get("l6_id") else None
    if status == "UNMATCHED":
        before_db = after_db = None

    update_db = None
    if saved:
        from src.database import get_progress_update
        stored = get_progress_update(saved["update_id"]) or {}
        update_db = {
            "update_id": saved["update_id"],
            "reported_by": stored.get("reported_by") or ingestion.get("reported_by", "Unknown Reporter"),
            "approved_by": stored.get("approved_by"),
            "match_status": stored.get("match_status") or saved["match_status"],
            "extraction_degraded": bool(stored.get("extraction_degraded") or match.get("extraction_degraded")),
            "duplicate": bool(saved["duplicate"]),
        }

    # Variance stays in whole days, whatever time of day was stated.
    variance_days = None
    finish_day = actual_day(after_db.get("actual_finish")) if after_db else None
    if finish_day and after_db.get("planned_finish"):
        try:
            variance_days = (finish_day - date.fromisoformat(after_db["planned_finish"])).days
        except ValueError:
            pass

    return {
        "event": event,
        "match": match,
        "before_db": before_db,
        "after_db": after_db,
        "update_db": update_db,
        "variance_days": variance_days,
        "outcome": outcome,
        "save_error": save_error,
    }


@app.post("/api/continue-to-match")
@role_required("contractor")
def continue_to_match():
    run = _current_run()
    if not run:
        return jsonify({"ok": False, "error": "No extracted results found."}), 400
    ingestion = run["ingestion"]
    if pipeline_runs.get_match_rows(run["run_id"]):
        # This run was already matched AND saved. A second click (or a
        # second tab) must not save again; just show the stored result.
        return jsonify({"ok": True, "redirect": url_for("matched"), "already_matched": True})
    items = pipeline_runs.list_items(run["run_id"])
    if not pipeline_runs.extraction_progress(items)["complete"]:
        return jsonify({"ok": False, "error": "Extraction is still running. Wait until every statement has finished."}), 409
    results = pipeline_runs.extracted_results(items)
    if not results:
        return jsonify({
            "ok": False,
            "error": "Nothing can be matched: every statement failed or needs correction. "
                     "Fix the statements on the ingestion screen and submit them again.",
        }), 400

    rows = []

    try:
        from src.database import get_planned_activity
    except Exception:
        get_planned_activity = None

    # Load the schedule ONCE for the whole batch (not once per event).
    schedule: list[dict[str, Any]] | None = None
    schedule_error: ScheduleUnavailableError | None = None
    try:
        schedule = get_match_schedule()
    except ScheduleUnavailableError as exc:
        app.logger.exception("Matching skipped for this batch: schedule unavailable")
        schedule_error = exc

    for event in results:
        if schedule is None:
            match = _schedule_unavailable_match(event, schedule_error)
        else:
            match = match_event(event, schedule)
        rows.append(_build_match_row(event, match, ingestion, get_planned_activity))

    pipeline_runs.save_match_rows(run["run_id"], rows)
    return jsonify({"ok": True, "redirect": url_for("matched")})


@app.post("/api/match/retry/<event_id>")
@role_required("contractor")
def retry_match(event_id: str):
    """Re-run the matcher for one event that ended in ERROR.

    Only ERROR rows can be retried. The result replaces that row in the run; it
    goes through the same guards as a first attempt.
    """
    run = _current_run()
    rows = pipeline_runs.get_match_rows(run["run_id"]) if run else None
    if not rows:
        return jsonify({"ok": False, "error": "No matched results found."}), 400
    idx = next((i for i, r in enumerate(rows) if (r.get("event") or {}).get("event_id") == event_id), None)
    if idx is None:
        return jsonify({"ok": False, "error": "That event is not part of the current run."}), 404
    if (rows[idx].get("match") or {}).get("review_status") != "ERROR":
        return jsonify({"ok": False, "error": "Only items that ended in an error can be retried."}), 409

    event = rows[idx]["event"]
    try:
        from src.database import get_planned_activity
    except Exception:
        get_planned_activity = None
    try:
        match = match_event(event, get_match_schedule())
    except ScheduleUnavailableError as exc:
        app.logger.exception("Retry skipped: schedule unavailable")
        match = _schedule_unavailable_match(event, exc)

    rows[idx] = _build_match_row(event, match, run["ingestion"], get_planned_activity)
    pipeline_runs.save_match_rows(run["run_id"], rows)
    return jsonify({"ok": True, "review_status": match.get("review_status"), "redirect": url_for("matched")})


@app.route("/matched")
@role_required("contractor")
def matched():
    run = _current_run()
    rows = pipeline_runs.get_match_rows(run["run_id"]) if run else None
    if not rows:
        return redirect(url_for("extracted"))
    return render_template(
        "matched.html",
        step=3,
        rows=rows,
        ingestion=run["ingestion"],
        **_stepper_availability(),
    )


@app.route("/memory")
@login_required
def memory():
    """Institutional Memory: queryable historical execution patterns.

    This page is READ-ONLY (it reads main_updates, matching_results and friends) -- it never writes to
    the database. Historical rows here come either from real approved
    updates, or from the optional `python -m src.seed_history` sample-data step.
    """
    from src.institutional_memory import (
        delay_reason_breakdown,
        discipline_productivity,
        distinct_filter_values,
        search_history,
        summary_stats,
    )

    filters = {
        "discipline": request.args.get("discipline") or None,
        "delay_category": request.args.get("delay_category") or None,
        "status": request.args.get("status") or None,
        "asset": request.args.get("asset") or None,
    }

    try:
        stats = summary_stats()
        productivity = discipline_productivity()
        delays = delay_reason_breakdown()
        filter_values = distinct_filter_values()
        history = search_history(**filters)
        error = None
    except Exception as exc:
        app.logger.exception("Institutional memory query failed")
        stats, productivity, delays, filter_values, history = {}, [], [], {"disciplines": [], "delay_categories": [], "statuses": []}, []
        error = str(exc)

    return render_template(
        "memory.html",
        step=0,
        stats=stats,
        productivity=productivity,
        delays=delays,
        filter_values=filter_values,
        history=history,
        filters=filters,
        error=error,
        hide_stepper=True,
    )


@app.post("/api/reset")
@role_required("contractor")
def reset():
    # Start over: discard only the current run's scratch data. The plan and its
    # history are not touched by this. Login is kept -- starting over shouldn't also
    # log the contractor out.
    run = _current_run()
    if run:
        pipeline_runs.delete_run(run["run_id"])
    session.pop("run_id", None)
    for key in _LEGACY_SESSION_KEYS:
        session.pop(key, None)
    return jsonify({"ok": True, "redirect": url_for("ingest")})


if __name__ == "__main__":
    # Warm the spaCy model and the retrieval index once at startup (~13s),
    # rather than paying that cost on whoever's first request happens to be
    # -- otherwise the first field statement anyone submits after starting
    # the server appears to hang for several seconds.
    try:
        print("Warming semantic matching engine (spaCy model + retrieval index)...")
        _warmup_schedule = get_match_schedule()
        from src.semantic_match import _get_index, _get_nlp
        _get_nlp()
        _get_index(_warmup_schedule)
        print("Ready.")
    except Exception:
        app.logger.exception("Warmup failed (non-fatal, will lazy-load on first request instead)")
    app.run(debug=_debug_enabled(), host="127.0.0.1", port=5000)
